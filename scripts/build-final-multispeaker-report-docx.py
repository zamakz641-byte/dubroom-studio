from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "DubRoom_Compte_Rendu_Complet_MultiSpeaker_GPT_SFX_2026-08-09.docx"
EXPORT_KIT = ROOT / "projects" / "sss-first-15min-v3" / "analysis" / "exchange" / "export-20260809-172248-manifest"
FINAL_EXPORT = ROOT / "exports" / "SSS Sacrificial Mage - test 15 min Multi-Speaker V3-20260809-170507"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
PALE_BLUE = "E8EEF5"
PALE_GREEN = "E7F4EA"
PALE_AMBER = "FFF3CD"
PALE_RED = "FDECEC"
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(92, 103, 117)
TABLE_WIDTH = 9360
TABLE_INDENT = 120


def set_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color if isinstance(color, RGBColor) else RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    if sum(widths) != TABLE_WIDTH:
        raise ValueError(f"La largeur du tableau doit totaliser {TABLE_WIDTH} DXA: {widths}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row_index, row in enumerate(table.rows):
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        tr_pr = row._tr.get_or_add_trPr()
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
        no_split = OxmlElement("w:cantSplit")
        tr_pr.append(no_split)
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=8, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])


def create_numbering(doc, bullet=False):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if bullet else "%1.")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr])
    if bullet:
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Calibri")
        fonts.set(qn("w:hAnsi"), "Calibri")
        r_pr.append(fonts)
        lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id):
    paragraph = doc.add_paragraph()
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.add_run(text)


def add_callout(doc, title, body, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    title_p = cell.paragraphs[0]
    title_p.paragraph_format.space_after = Pt(3)
    set_font(title_p.add_run(title), size=10.5, color=DARK_BLUE, bold=True)
    body_p = cell.add_paragraph(body)
    body_p.paragraph_format.space_after = Pt(0)
    body_p.paragraph_format.line_spacing = 1.15
    set_table_geometry(table, [TABLE_WIDTH])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths, path_columns=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(label), size=8.5, color=DARK_BLUE, bold=True)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_font(p.add_run(str(value)), name="Consolas" if index in path_columns else "Calibri", size=7.3 if index in path_columns else 8.25, color=INK)
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    header_table = section.header.add_table(rows=1, cols=2, width=Inches(6.5))
    left = header_table.cell(0, 0).paragraphs[0]
    right = header_table.cell(0, 1).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    right.paragraph_format.space_after = Pt(0)
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(left.add_run("DUBROOM  /  RAPPORT D’INTÉGRATION"), size=8, color=DARK_BLUE, bold=True)
    set_font(right.add_run("09 AOÛT 2026"), size=8, color=MUTED)
    set_table_geometry(header_table, [6500, 2860])
    add_page_field(section.footer.paragraphs[0])


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def build():
    doc = Document()
    configure(doc)
    bullet_id = create_numbering(doc, bullet=True)
    normal_number_id = create_numbering(doc, bullet=False)
    multi_number_id = create_numbering(doc, bullet=False)
    next_test_number_id = create_numbering(doc, bullet=False)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(4)
    set_font(title.add_run("DubRoom — compte rendu complet"), size=25, color=DARK_BLUE, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    set_font(subtitle.add_run("Multi-Speaker, GPT R3, voix OmniVoice, synchronisation et SFX préservés"), size=13, color=MUTED)
    meta = [
        ["Application", str(ROOT)],
        ["Projet de référence", "sss-first-15min-v3 — test réel d’environ 15 minutes"],
        ["Révision GPT active", "narrator-speaker-sfx-safe-20260809-r3"],
        ["État", "Fonctions intégrées et tests locaux réussis ; nouveau script GPT à régénérer"],
    ]
    add_table(doc, ["Repère", "Valeur"], meta, [2700, 6660], path_columns=(1,))

    heading(doc, "Résumé exécutif")
    add_callout(
        doc,
        "Résultat",
        "Le parcours principal peut maintenant être piloté depuis DubRoom : choix Normal ou Multi-Speaker, analyse audio, transcription chinoise spécialisée, indices de locuteur, export du Kit GPT, validation du JSON réimporté, casting des voix, génération OmniVoice sur CUDA, synchronisation et export avec musique/SFX conservés.",
        PALE_GREEN,
    )
    for text in (
        "Le mode Normal reste sur le manifeste V2 ; le mode Multi-Speaker utilise exclusivement le manifeste V3.",
        "Le GPT devient le résolveur central des personnages grâce au texte, au contexte et aux indices acoustiques, sans transformer chaque changement de timbre en nouveau personnage.",
        "Le mix final conserve désormais le lit d’ambiance même lorsque l’ancienne voix d’origine est coupée.",
        "Le nouveau prompt GPT impose les personnages stables, les émotions et événements OmniVoice, les sections de locuteurs et un fichier JSON téléchargeable.",
    ):
        add_list_item(doc, text, bullet_id)

    doc.add_page_break()
    heading(doc, "1. Parcours désormais disponible dans l’application")
    heading(doc, "Mode Normal — une seule voix", 2)
    for text in (
        "Créer le projet et choisir Normal.",
        "Analyser/transcrire la source, puis préparer le Kit GPT au format V2.",
        "Réimporter la traduction, choisir une voix unique, générer, synchroniser et exporter.",
    ):
        add_list_item(doc, text, normal_number_id)
    heading(doc, "Mode Multi-Speaker — narrateur et personnages", 2)
    for text in (
        "Créer le projet et choisir Multi-Speaker dès le départ.",
        "Préserver l’audio et transcrire le chinois avec FunASR Paraformer-zh lorsque la langue le demande.",
        "Produire les indices techniques de changement de voix ; ils restent des indices, jamais l’identité finale.",
        "Exporter le Kit GPT V3. Le GPT traduit et résout les personnages par section avec le contexte narratif.",
        "Réimporter le JSON. DubRoom bloque un manifeste incomplet ou incohérent avant qu’il ne contamine le casting.",
        "Attribuer/cloner les voix dans la bibliothèque, générer les unités vocales avec émotions, synchroniser et exporter.",
    ):
        add_list_item(doc, text, multi_number_id)
    add_callout(doc, "Règle narrative active", "NARRATOR raconte et complète les temps morts autorisés. MC possède sa propre voix lorsqu’il parle réellement ; les pensées privées peuvent rester MC/inner_monologue. Les autres personnages conservent chacun leur propre voice_key stable.")

    doc.add_page_break()
    heading(doc, "2. Nouveau contrat GPT R3")
    add_table(
        doc,
        ["Contrôle", "Règle active"],
        [
            ["Sortie", "Fichier JSON téléchargeable <project_id>-completed-v3.json ; texte brut .txt accepté seulement en secours."],
            ["Personnages", "Registre stable par section : speaker_id, voice_key, sexe vocal, tranche d’âge, rôle et confiance."],
            ["Narrateur", "N’utilise que narration/bridge ; ne vole pas les répliques prouvées du MC ou des autres personnages."],
            ["Unités", "IDs séquentiels <segment>-u01, ordre strict, pas de fragment pendant ou isolé."],
            ["OmniVoice", "Emotion + événements contextuels tels que [laugh], [sigh], [gasp], [cry], [whisper]."],
            ["SFX", "Pont narrateur autorisé uniquement quand available=true et sfx_safe=true."],
            ["Validation", "Le GPT doit vérifier programmatiquement le JSON avant de fournir le fichier."],
        ],
        [2350, 7010],
    )
    add_callout(doc, "Limite respectée", "Le fichier d’instructions actif contient 7 550 caractères, donc reste sous la limite de 8 000 caractères demandée.", PALE_GREEN)

    doc.add_page_break()
    heading(doc, "3. Import, casting et bibliothèque vocale")
    for text in (
        "Le parseur accepte un vrai .json, un JSON brut enregistré en .txt et un bloc ```json ...``` copié depuis ChatGPT.",
        "La prévisualisation signale les personnages invalides, unités désordonnées, émotions inconnues, bridges dangereux, sections incomplètes, registre manquant et ancienne révision de prompt.",
        "Un projet Multi-Speaker exige maintenant un manifeste complet avant import définitif.",
        "La bibliothèque classe les profils par moteur, langue, sexe vocal, tranche d’âge, rôle narratif et usage préféré (MC, héroïne principale, antagoniste, narrateur ou secondaire).",
        "OmniVoice sert au Multi-Speaker et prend en charge clonage, voice design, enregistrement et événements expressifs ; les autres moteurs restent classés comme alternatives ou Single-Speaker selon leurs capacités.",
    ):
        add_list_item(doc, text, bullet_id)
    add_callout(doc, "OmniVoice réparé", "La génération réelle a terminé sur CUDA sans l’erreur WeTextProcessing/tn. Le warning expandable_segments peut encore apparaître sous Windows, mais il est sans effet fonctionnel.", PALE_GREEN)

    heading(doc, "4. Conservation audio, SFX et export")
    add_table(
        doc,
        ["Avant", "Après correction"],
        [
            ["original_volume=0 pouvait supprimer tout le lit audio.", "La voix source est coupée, mais bed.wav reste présent avec musique et SFX."],
            ["Le volume d’ambiance suivait le mauvais réglage.", "bed.wav utilise le réglage Musique/SFX préservés, 0,85 par défaut."],
            ["Le ducking risquait d’écraser les effets.", "Ducking désactivé par défaut ; activation volontaire seulement."],
            ["Concaténation FLAC corrompue après environ 6 min 05 s.", "Retiming et master interne en PCM 24 bits, concat vidéo sûre puis encodage final."],
            ["Contrôles audio peu clairs.", "Interface renommée explicitement Musique / SFX préservés."],
        ],
        [4300, 5060],
    )
    add_callout(doc, "Limite actuelle", "Le projet de 15 minutes possède encore un bed.wav combiné. Les SFX sont conservés, mais Bandit peut laisser du dialogue résiduel ou altérer certains effets. Des stems music.wav et sfx.wav indépendants donneraient un contrôle supérieur lors d’une future amélioration.", PALE_AMBER)

    doc.add_page_break()
    heading(doc, "5. Moteurs, GPU et dépendances communes")
    add_table(
        doc,
        ["Composant", "État", "Détail"],
        [
            ["OmniVoice HQ", "Prêt", "CUDA réel ; génération testée en 20,06 s modèle inclus pour 1,97 s d’audio."],
            ["FunASR zh", "Prêt", "funasr 1.2.7, Paraformer-zh + FSMN-VAD, auto-test CUDA réussi."],
            ["Demucs", "Prêt", "Profil rapide GPU ; htdemucs standard recommandé plutôt que htdemucs_ft."],
            ["Bandit DnR3", "Prêt", "Profil qualité cinéma ; meilleure conservation sémantique de l’ambiance que Demucs sur le test réel."],
            ["Sherpa/TitaNet", "Prêt", "Indices de tours de parole locaux ; le GPT décide ensuite de l’identité narrative."],
            ["Planificateur GPU", "Intégré", "Évite les téléchargements répétés et centralise les dépendances ; charge longue simultanée à éprouver."],
        ],
        [2200, 1300, 5860],
    )
    heading(doc, "Principe de dépendances partagées", 2)
    for text in (
        "Réutiliser les environnements et modèles déjà présents avant toute installation.",
        "Conserver un ready.json et une empreinte pour savoir si un moteur est réellement prêt.",
        "Éviter plusieurs copies de PyTorch/CUDA pour les moteurs compatibles.",
        "Mettre en cache les résultats lourds et ne retraiter que si la source ou les paramètres changent.",
    ):
        add_list_item(doc, text, bullet_id)

    heading(doc, "6. Preuves de validation")
    checks = [
        ["Interface desktop", "Réussi", "596 clés i18n, TypeScript et build Vite validés."],
        ["Échange GPT V2/V3", "Réussi", "1 160 segments, round-trip V3, V2 Single, JSON/.txt/fenced JSON."],
        ["Export audio", "Réussi", "SFX audibles entre les prises de voix ; marqueur bed interne conservé."],
        ["Master audio", "Réussi", "48 kHz stéréo PCM 24 bits et filtres validés."],
        ["Multi-Speaker", "Réussi", "Pipeline, persistance, diarisation et métadonnées vocales validés."],
        ["Bibliothèque vocale", "Réussi", "18 profils, mode local uniquement et CUDA détecté."],
        ["TTS natifs", "Réussi", "OmniVoice inclus dans le catalogue et les adaptateurs."],
        ["Pipeline ouvert", "Réussi", "Blocs naturels fusionnés au lieu de micro-segments artificiels."],
        ["Export vidéo 15 min", "Réussi", "924,233 s, décodage complet, bed mesuré à 30 s, 450 s et 810 s."],
    ]
    add_table(doc, ["Contrôle", "État", "Preuve"], checks, [2600, 1400, 5360])

    doc.add_page_break()
    heading(doc, "7. Nouveaux fichiers principaux")
    new_files = [
        ["services/api/app/audio_preservation_service.py", "Préservation et séparation audio."],
        ["services/api/app/diarization_service.py", "Indices de locuteur et profils vocaux."],
        ["services/api/app/native_tts_service.py", "Orchestration des TTS natifs."],
        ["services/api/app/resource_scheduler.py", "Arbitrage des ressources lourdes/GPU."],
        ["services/api/app/shared_dependency_service.py", "Réutilisation des dépendances communes."],
        ["services/api/app/sync_service.py", "Synchronisation continue et concat PCM sûre."],
        ["services/api/app/transcript_exchange_service.py", "Manifestes V2/V3, Kit GPT et validation d’import."],
        ["services/api/app/tts_catalog_service.py", "Catalogue et capacités des moteurs TTS."],
        ["services/api/app/voice_library_service.py", "Bibliothèque de voix et métadonnées."],
        ["services/api/app/voice_reference_service.py", "Références, clonage et enregistrement."],
        ["apps/desktop/src/components/YouTubeRangeSelector.tsx", "Sélection d’un extrait YouTube."],
        ["apps/desktop/src/components/studio/AudioPreservationPanel.tsx", "Choix qualité/rapide et état audio."],
        ["apps/desktop/src/components/studio/SpeakerDiarizationPanel.tsx", "Détection et validation des locuteurs."],
        ["apps/desktop/src/components/voice/OmniVoiceLibraryPanel.tsx", "Interface dédiée OmniVoice."],
        ["scripts/engines/run-funasr-zh.py", "Transcription chinoise spécialisée."],
        ["scripts/engines/run-sherpa-diarization.py", "Worker local de diarisation."],
        ["scripts/engines/run-audio-separation.py", "Worker de séparation audio."],
        ["scripts/engines/run-cinematic-separation.py", "Séparation Bandit cinéma."],
        ["scripts/engines/run-tts.py", "Adaptateurs TTS et OmniVoice."],
        ["docs/custom-gpt-dubroom/README_ACTIF.md", "Carte des instructions GPT réellement actives."],
    ]
    add_table(doc, ["Fichier", "Rôle"], new_files, [6000, 3360], path_columns=(0,))

    doc.add_page_break()
    heading(doc, "8. Fichiers modifiés principaux")
    modified_files = [
        ["services/api/app/main.py", "Modes projet, endpoints GPT, casting, mix et export."],
        ["services/api/app/export_service.py", "Bed audio, volumes, ducking et PCM 24 bits."],
        ["services/api/app/engine_service.py", "Installation et disponibilité des moteurs."],
        ["services/api/app/local_voice_service.py", "Voix locales et compatibilités."],
        ["services/api/app/youtube_service.py", "Import et plages YouTube."],
        ["apps/desktop/src/pages/StudioPage.tsx", "Kit GPT, erreurs V3, mix, sync et export."],
        ["apps/desktop/src/pages/LibraryPage.tsx", "Classement et gestion de la bibliothèque."],
        ["apps/desktop/src/components/voice/ProfileCreator.tsx", "Clonage, design, enregistrement et rôles."],
        ["apps/desktop/src/i18n.tsx", "Libellés français/anglais et messages de validation."],
        ["apps/desktop/src/types.ts", "Types des manifestes, locuteurs et aperçus."],
        ["apps/desktop/src/lib/api.ts", "Contrats API de l’interface."],
        ["config/app.config.json", "Moteurs, options et politiques locales."],
        ["package.json", "Commandes de build et tests."],
        ["scripts/engines/run-asr.py", "Routage ASR et moteur chinois."],
        ["scripts/engines/install-python-engine.ps1", "Installation mutualisée des moteurs Python."],
    ]
    add_table(doc, ["Fichier", "Modification"], modified_files, [5900, 3460], path_columns=(0,))

    heading(doc, "Documentation GPT active", 2)
    gpt_files = [
        ["docs/custom-gpt-dubroom/INSTRUCTIONS_CHINESE_MULTISPEAKER_V3_8000.txt", "Instructions R3 à coller dans le GPT ; 7 550 caractères."],
        ["docs/custom-gpt-dubroom/KNOWLEDGE_MULTISPEAKER_JSON_V3.txt", "Base de connaissances du format Multi-Speaker V3."],
        ["docs/custom-gpt-dubroom/CREER_LE_GPT.md", "Procédure de configuration et fichiers à charger."],
        [str(EXPORT_KIT.relative_to(ROOT) / "sss-first-15min-v3-translation-manifest.json"), "Manifeste neuf exporté par l’application."],
    ]
    add_table(doc, ["Emplacement", "Usage"], gpt_files, [5900, 3460], path_columns=(0,))

    doc.add_page_break()
    heading(doc, "9. Problèmes encore présents")
    issues = [
        ["P1", "Ancien script à refaire", "Le manifeste actuel vient du prompt R2. Il reste compatible, mais doit être régénéré avec le Kit R3 avant le prochain rendu."],
        ["P1", "Qualité du script existant", "4 unités UNKNOWN et 15 unités de deux mots ou moins ; fragments notables : “She’s only” et “Qingshan Guild...”"],
        ["P1", "Étape GPT externe", "L’application prépare et valide tout, mais le Custom GPT privé reste manuel : exporter, joindre, télécharger, réimporter."],
        ["P2", "Séparation SFX", "bed.wav est combiné ; certains effets peuvent encore être modifiés ou contenir du bleed vocal."],
        ["P2", "ASR chinois", "FunASR réduit les phrases corrompues, mais noms propres et voix couvertes par les SFX doivent être revus sur une vraie vidéo."],
        ["P2", "Sur-segmentation acoustique", "Le test avait 87 clusters temporaires. Le GPT doit les fusionner en personnages stables."],
        ["P2", "Concurrence GPU", "Le planificateur est intégré, mais un run long avec plusieurs moteurs lourds n’est pas encore stress-testé."],
        ["P3", "Warning Windows", "expandable_segments not supported peut rester visible ; il n’empêche pas OmniVoice de fonctionner."],
        ["P3", "Taille de l’interface", "Avertissement de bundle supérieur à 500 kB ; aucun impact fonctionnel, découpage futur possible."],
    ]
    add_table(doc, ["Priorité", "Sujet", "État / action"], issues, [1100, 2400, 5860])
    add_callout(doc, "Redémarrage requis", "Si DubRoom ou son service local était déjà ouvert pendant les modifications, fermer puis relancer l’application afin de charger le nouveau code.", PALE_AMBER)

    doc.add_page_break()
    heading(doc, "10. Prochain test recommandé")
    for text in (
        "Relancer DubRoom et ouvrir le projet Multi-Speaker.",
        "Exporter le nouveau Kit GPT daté du 9 août 2026.",
        "Configurer le Custom GPT avec les instructions R3 et la base de connaissances V3 actives.",
        "Régénérer le script des 15 minutes ; vérifier personnages, narrator/MC, émotions et absence de fragments.",
        "Réimporter : ne continuer que si la prévisualisation ne signale aucune erreur bloquante.",
        "Caster une voix marquante pour NARRATOR/MC, l’héroïne et l’antagoniste, puis générer sur OmniVoice CUDA.",
        "Synchroniser, écouter des points riches en SFX et exporter la vidéo finale.",
    ):
        add_list_item(doc, text, next_test_number_id)
    add_callout(doc, "Critère de réussite", "Une vidéo de 15 à 20 minutes où le narrateur reste distinct du MC, chaque personnage garde sa voix, les événements expressifs sont naturels, aucun dialogue n’est muet et les SFX restent audibles.", PALE_GREEN)

    heading(doc, "11. Emplacements de livraison")
    delivery = [
        [str(OUTPUT), "Présent compte rendu Word."],
        [str(EXPORT_KIT), "Kit GPT R3 fraîchement produit par DubRoom."],
        [str(FINAL_EXPORT), "Dernier export vidéo 15 min avec SFX rétablis."],
        [str(ROOT / "data" / "environments" / "asr-funasr-zh" / "ready.json"), "Preuve d’installation FunASR CUDA."],
        [str(ROOT / "data" / "tts" / "audio" / "1df19e8a4a174015a294490013315ac8.wav"), "Échantillon de génération OmniVoice CUDA."],
    ]
    add_table(doc, ["Chemin", "Contenu"], delivery, [6500, 2860], path_columns=(0,))

    doc.core_properties.title = "DubRoom — compte rendu complet Multi-Speaker, GPT et SFX"
    doc.core_properties.subject = "État de l’intégration au 9 août 2026"
    doc.core_properties.author = "DubRoom / Codex"
    doc.core_properties.keywords = "DubRoom, Multi-Speaker, GPT R3, OmniVoice, FunASR, SFX, synchronisation"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(json.dumps({"output": str(OUTPUT), "bytes": OUTPUT.stat().st_size}, ensure_ascii=False))


if __name__ == "__main__":
    build()
